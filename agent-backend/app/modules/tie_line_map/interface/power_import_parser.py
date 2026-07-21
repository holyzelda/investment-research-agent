"""联络线功率 Excel 导入解析。"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile

import openpyxl
from fastapi import UploadFile
from openpyxl.utils.exceptions import InvalidFileException

from app.core.exceptions import ValidationException
from app.modules.tie_line_map.application.power_import_service import (
    POWER_TIME_COLUMNS,
    PowerImportRecord,
)

_EXPECTED_HEADER = ("日期", "名称", "电压等级", *POWER_TIME_COLUMNS)


def _parse_date(value: object, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value)
        except ValueError as error:
            raise ValidationException(f"Excel 第 {row_number} 行日期格式无效") from error
    raise ValidationException(f"Excel 第 {row_number} 行日期格式无效")


def _parse_power(value: object, row_number: int, time_column: str) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        power = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise ValidationException(
            f"Excel 第 {row_number} 行 {time_column} 功率必须为数值"
        ) from error
    if not power.is_finite():
        raise ValidationException(f"Excel 第 {row_number} 行 {time_column} 功率必须为有限数值")
    return power


async def parse_excel_import(upload: UploadFile) -> list[PowerImportRecord]:
    if not upload.filename or not upload.filename.lower().endswith(".xlsx"):
        raise ValidationException("仅支持 .xlsx 文件")
    content = await upload.read()
    if not content:
        raise ValidationException("Excel 文件不能为空")
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ValidationException("Excel 文件格式无效") from error
    try:
        if "功率" not in workbook.sheetnames:
            raise ValidationException('Excel 缺少“功率”工作表')
        sheet = workbook["功率"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if tuple(header or ()) != _EXPECTED_HEADER:
            raise ValidationException("Excel 表头与功率模板不一致")
        records = []
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            if len(row) != len(_EXPECTED_HEADER):
                raise ValidationException(f"Excel 第 {row_number} 行列数不正确")
            data_date = _parse_date(row[0], row_number)
            line_name = str(row[1] or "").strip()
            voltage_text = str(row[2] or "").strip()
            if not line_name or not voltage_text:
                raise ValidationException(f"Excel 第 {row_number} 行名称和电压等级不能为空")
            powers = {
                time_column: _parse_power(value, row_number, time_column)
                for time_column, value in zip(POWER_TIME_COLUMNS, row[3:])
            }
            records.append(PowerImportRecord(data_date, line_name, voltage_text, powers))
        if not records:
            raise ValidationException("Excel 不包含功率数据")
        return records
    finally:
        workbook.close()
