"""上传现货电价的解析与校验。"""
from __future__ import annotations

import hashlib
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook

from app.core.exceptions import ValidationException

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_ROWS = 1_000_000
MAX_DATE_SPAN_DAYS = 366
PRICE_UNIT = "CNY/MWh"
ALGORITHM_VERSION = "spot-dp-v1"
EXPECTED_HEADERS = ("日期", "时点", "电价（元/MWh）")
DATA_SHEET_NAME = "现货电价"
INSTRUCTION_SHEET_NAME = "填写说明"
_GRID_LABELS = {
    f"{(index * 15) // 60:02d}:{(index * 15) % 60:02d}"
    for index in range(1, 97)
}


@dataclass(frozen=True)
class ParsedSpotContext:
    file_name: str
    file_sha256: str
    start_date: date
    end_date: date
    valid_day_count: int
    skipped_day_count: int
    warnings: list[str]
    k_results: dict[str, dict[str, Any]]
    raw_prices: dict[str, list[str]]
    price_mode: str = "both"


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_time(value: object) -> str | None:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            return None
        if value.hour == 0 and value.minute == 0:
            return "24:00"
        label = f"{value.hour:02d}:{value.minute:02d}"
        return label if label in _GRID_LABELS else None
    if isinstance(value, (int, float)) and 0 <= value <= 1:
        minutes = round(float(value) * 24 * 60)
        if minutes % 15 == 0 and 0 < minutes <= 24 * 60:
            return f"{minutes // 60:02d}:{minutes % 60:02d}"
        return None
    text = str(value or "").strip()
    if text in {"24:00", "24:00:00"}:
        return "24:00"
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            parsed = datetime.strptime(text, fmt).time()
        except ValueError:
            continue
        if parsed.second:
            return None
        label = f"{parsed.hour:02d}:{parsed.minute:02d}"
        return label if label in _GRID_LABELS else None
    return None


def _parse_price(value: object) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        parsed = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None
    return parsed if parsed.is_finite() else None


def _error(message: str) -> ValidationException:
    return ValidationException(f"电价文件校验失败：{message}")


def parse_uploaded_spot_file(
    content: bytes,
    filename: str,
    *,
    min_action_minutes: int = 30,
    switch_gap_minutes: int = 15,
    workers: int | None = None,
    price_mode: str = "both",
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
) -> ParsedSpotContext:
    """解析标准长表并保存完整日的原始 96 点曲线；DP 延迟到实际测算时执行。"""
    if price_mode not in {"both", "charge_only", "discharge_only"}:
        raise _error("电价用途无效，必须为 both、charge_only 或 discharge_only")
    if not filename or not filename.lower().endswith(".xlsx"):
        raise _error("仅支持 .xlsx 文件")
    if not content:
        raise _error("文件为空")
    if len(content) > MAX_UPLOAD_BYTES:
        raise _error(f"文件不能超过 {MAX_UPLOAD_BYTES // (1024 * 1024)}MB")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError) as exc:
        raise _error("Excel 文件无法打开") from exc

    try:
        if progress_callback is not None:
            progress_callback({"phase": "validating", "message": "正在读取并校验 Excel 数据"})
        sheet_names = workbook.sheetnames
        if DATA_SHEET_NAME in sheet_names:
            unexpected_sheets = [
                name for name in sheet_names
                if name not in {DATA_SHEET_NAME, INSTRUCTION_SHEET_NAME}
            ]
            if unexpected_sheets:
                raise _error("除现货电价和填写说明外，不能包含其他工作表")
            sheet = workbook[DATA_SHEET_NAME]
        elif len(sheet_names) == 1:
            # 兼容本功能上线前生成的单工作表模板和用户已有文件。
            sheet = workbook.active
        else:
            raise _error("包含多个工作表时，数据工作表必须命名为“现货电价”")
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in (next(rows, None) or ()))
        if headers != EXPECTED_HEADERS:
            expected = "、".join(EXPECTED_HEADERS)
            raise _error(f"表头不匹配，必须为：{expected}")

        by_day: dict[date, dict[str, Decimal]] = defaultdict(dict)
        invalid_days: set[date] = set()
        malformed_rows: list[str] = []
        row_count = 1
        for row_number, row in enumerate(rows, start=2):
            row_count += 1
            if row_count > MAX_ROWS:
                raise _error(f"数据行不能超过 {MAX_ROWS:,} 行")
            if not any(value not in (None, "") for value in row):
                continue
            if len(row) < 3:
                raise _error(f"第 {row_number} 行列数不足")
            business_date = _parse_date(row[0])
            if business_date is None:
                malformed_rows.append(f"第{row_number}行日期无效")
                continue
            label = _parse_time(row[1])
            price = _parse_price(row[2])
            if label is None or price is None:
                invalid_days.add(business_date)
                continue
            if label in by_day[business_date]:
                invalid_days.add(business_date)
            else:
                by_day[business_date][label] = price

        if malformed_rows:
            detail = "、".join(malformed_rows[:3])
            suffix = "等" if len(malformed_rows) > 3 else ""
            raise _error(f"{detail}{suffix}，无法识别日期")

        valid_days: dict[date, list[Decimal]] = {}
        for business_date, cells in by_day.items():
            if business_date in invalid_days or len(cells) != 96:
                invalid_days.add(business_date)
                continue
            ordered = [cells.get(
                f"{(index * 15) // 60:02d}:{(index * 15) % 60:02d}"
            ) for index in range(1, 97)]
            if any(price is None for price in ordered):
                invalid_days.add(business_date)
                continue
            valid_days[business_date] = ordered  # type: ignore[assignment]

        if not valid_days:
            raise _error("没有包含完整的 96 点有效自然日")
        first_day = min(by_day)
        last_day = max(by_day)
        if (last_day - first_day).days + 1 > MAX_DATE_SPAN_DAYS:
            raise _error(f"数据日期范围不能超过 {MAX_DATE_SPAN_DAYS} 天")

        skipped_count = (last_day - first_day).days + 1 - len(valid_days)
        warnings = []
        if skipped_count:
            warnings.append(f"有{skipped_count}个自然日因时点不完整、重复或价格无效未参与计算")
        return ParsedSpotContext(
            file_name=filename[:255],
            file_sha256=hashlib.sha256(content).hexdigest(),
            start_date=first_day,
            end_date=last_day,
            valid_day_count=len(valid_days),
            skipped_day_count=skipped_count,
            warnings=warnings,
            raw_prices={
                business_date.isoformat(): [str(price) for price in prices]
                for business_date, prices in valid_days.items()
            },
            # 保留旧字段以兼容历史 context；新上传不在此阶段计算 DP。
            k_results={},
            price_mode=price_mode,
        )
    finally:
        workbook.close()
