#!/usr/bin/env python3
"""将栏目抓取生成的 Excel 待入库数据写入 MySQL。"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

try:
    from .database_config import load_database_config
except ImportError:
    from database_config import load_database_config


TABLE_NAME = "policy_market_scrolling_messages"
IMPORT_HEADERS = (
    "title",
    "summary",
    "link_url",
    "province",
    "publish_time",
    "expire_time",
    "status",
    "rating",
    "sort_order",
    "tags",
    "extra_data",
    "is_deleted",
    "created_by",
    "updated_by",
)
INSERT_COLUMNS = IMPORT_HEADERS
PROVINCE_ALIASES = {
    "内蒙古自治区": "内蒙古",
    "广西壮族自治区": "广西",
    "西藏自治区": "西藏",
    "宁夏回族自治区": "宁夏",
    "新疆维吾尔自治区": "新疆",
    "香港特别行政区": "香港",
    "澳门特别行政区": "澳门",
}


def parse_args() -> argparse.Namespace:
    yesterday = date.today() - timedelta(days=1)
    parser = argparse.ArgumentParser(
        description="将北极星栏目 Excel 数据写入 policy_market_scrolling_messages。"
    )
    parser.add_argument(
        "--date",
        type=lambda value: datetime.strptime(value, "%Y-%m-%d").date(),
        default=yesterday,
        help="默认读取昨日生成的 Excel，格式 YYYY-MM-DD。",
    )
    parser.add_argument(
        "--input",
        nargs="+",
        type=Path,
        help="指定一个或多个 Excel 文件；不指定时自动读取当前目录下对应日期的北极星 Excel。",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="只校验并统计，不连接数据库、不写入数据。",
    )
    return parser.parse_args()


def find_input_files(input_paths: list[Path] | None, target_date: date) -> list[Path]:
    if input_paths:
        files = input_paths
    else:
        pattern = f"北极星*_{target_date.isoformat()}.xlsx"
        files = sorted(Path.cwd().glob(pattern))
    files = [path.resolve() for path in files]
    missing = [str(path) for path in files if not path.is_file()]
    if missing:
        raise RuntimeError(f"Excel 文件不存在：{', '.join(missing)}")
    if not files:
        raise RuntimeError(
            f"未找到 Excel 文件，请检查当前目录，或使用 --input 指定文件。日期：{target_date.isoformat()}"
        )
    return files


def as_datetime(value: object, field: str, row_number: int, path: Path) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                continue
    raise RuntimeError(f"{path.name} 第 {row_number} 行的 {field} 不是有效日期时间。")


def normalize_province(value: object) -> str | None:
    """将省级行政区全称统一为数据库使用的简称。"""
    if value in (None, ""):
        return None
    province = str(value).strip()
    if not province:
        return None
    if province in PROVINCE_ALIASES:
        return PROVINCE_ALIASES[province]
    if province.endswith(("省", "市")):
        return province[:-1]
    return province


def read_excel_records(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        headers = tuple(cell.value for cell in next(worksheet.iter_rows(max_row=1)))
        if headers != IMPORT_HEADERS:
            raise RuntimeError(
                f"{path.name} 表头不匹配，必须是数据库待入库字段顺序：{', '.join(IMPORT_HEADERS)}"
            )

        records = []
        for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(value not in (None, "") for value in values):
                continue
            record = dict(zip(IMPORT_HEADERS, values))
            title = str(record["title"] or "").strip()
            link_url = str(record["link_url"] or "").strip()
            if not title or len(title) > 255:
                raise RuntimeError(f"{path.name} 第 {row_number} 行的 title 为空或超过 255 字符。")
            if not link_url or len(link_url) > 1000:
                raise RuntimeError(f"{path.name} 第 {row_number} 行的 link_url 为空或超过 1000 字符。")
            record["title"] = title
            record["link_url"] = link_url
            record["province"] = normalize_province(record["province"])
            record["publish_time"] = as_datetime(record["publish_time"], "publish_time", row_number, path)
            record["expire_time"] = as_datetime(record["expire_time"], "expire_time", row_number, path)
            try:
                json.loads(str(record["extra_data"] or "{}"))
            except json.JSONDecodeError as exc:
                raise RuntimeError(f"{path.name} 第 {row_number} 行的 extra_data 不是合法 JSON。") from exc
            records.append(record)
        return records
    finally:
        workbook.close()


def write_records(
    records: list[dict[str, object]], database_config: dict[str, object]
) -> tuple[int, int]:
    try:
        import pymysql
    except ImportError as exc:
        raise RuntimeError("缺少 PyMySQL，请先执行：python -m pip install -r requirements.txt") from exc

    columns = ", ".join(f"`{column}`" for column in INSERT_COLUMNS)
    placeholders = ", ".join(["%s"] * len(INSERT_COLUMNS))
    insert_sql = f"INSERT INTO `{TABLE_NAME}` ({columns}) VALUES ({placeholders})"
    select_sql = f"SELECT id FROM `{TABLE_NAME}` WHERE link_url = %s AND is_deleted = 0 LIMIT 1"
    connection = None
    inserted = skipped = 0
    try:
        connection = pymysql.connect(**database_config)
        with connection.cursor() as cursor:
            for record in records:
                cursor.execute(select_sql, (record["link_url"],))
                if cursor.fetchone():
                    skipped += 1
                    continue
                cursor.execute(insert_sql, tuple(record[column] for column in INSERT_COLUMNS))
                inserted += 1
        connection.commit()
    except Exception as exc:
        if connection is not None:
            connection.rollback()
        raise RuntimeError(f"数据库写入失败，已回滚：{exc}") from exc
    finally:
        if connection is not None:
            connection.close()
    return inserted, skipped


def main() -> int:
    args = parse_args()
    try:
        files = find_input_files(args.input, args.date)
        records = []
        for path in files:
            file_records = read_excel_records(path)
            print(f"已读取：{path.name}，{len(file_records)} 条")
            records.extend(file_records)
        unique_records = []
        seen_urls = set()
        for record in records:
            if record["link_url"] in seen_urls:
                continue
            seen_urls.add(record["link_url"])
            unique_records.append(record)
        print(f"读取总数：{len(records)}，文件内合并去重后：{len(unique_records)}")
        if args.dry_run:
            print("预览模式：未连接数据库，未写入数据。")
            return 0
        database_config = load_database_config()
        inserted, skipped = write_records(unique_records, database_config)
    except RuntimeError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    print(f"数据库新增：{inserted}，数据库已有跳过：{skipped}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
