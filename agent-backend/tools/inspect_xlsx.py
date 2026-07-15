"""一次性探查：列出三个 xlsx 的 sheet 结构，预览映射表"副本"sheet 表头与前几行。"""
import openpyxl

BASE = r"D:/0_工作/2026/投研智能体/测算表"
FILES = {
    "映射表": rf"{BASE}/Excel单元格映射表_业务填写清单_20260630.xlsx",
    "表2": rf"{BASE}/各省电力市场及收益结构1.xlsx",
    "引擎表": rf"{BASE}/15年不换电-表前储能经评模型V5.0.xlsx",
}


def sheet_inventory():
    for label, path in FILES.items():
        print(f"\n===== {label}: {path.split('/')[-1]} =====")
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            print(f"  打开失败: {e}")
            continue
        for ws in wb.worksheets:
            print(f"  - sheet: 「{ws.title}」  dims={ws.max_row}行 x {ws.max_column}列")
        wb.close()


def preview_mapping_copy():
    path = FILES["映射表"]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # 找"副本"sheet
    target = None
    for ws in wb.worksheets:
        if "副本" in ws.title:
            target = ws
            break
    if target is None:
        print("\n[映射表] 未找到含'副本'的 sheet，列出所有 sheet：", [w.title for w in wb.worksheets])
        wb.close()
        return
    print(f"\n===== 映射表「{target.title}」前 12 行预览（A..L 列）=====")
    for i, row in enumerate(target.iter_rows(min_row=1, max_row=12, max_col=12, values_only=True), start=1):
        cells = [("" if v is None else str(v))[:18] for v in row]
        print(f"  r{i:>2}: " + " | ".join(cells))
    wb.close()


if __name__ == "__main__":
    sheet_inventory()
    preview_mapping_copy()
