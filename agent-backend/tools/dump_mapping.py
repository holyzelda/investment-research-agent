"""提取映射表「副本」全部参数行 + 表2 各省收益结构表头，作为逐项评审底稿。"""
import openpyxl

BASE = r"D:/0_工作/2026/投研智能体/测算表"
MAP = rf"{BASE}/Excel单元格映射表_业务填写清单_20260630.xlsx"
TBL2 = rf"{BASE}/各省电力市场及收益结构1.xlsx"


def dump_mapping():
    wb = openpyxl.load_workbook(MAP, read_only=True, data_only=True)
    ws = next(w for w in wb.worksheets if "副本" in w.title)
    # 列：A序号 B类别 C所属sheet D单元格 E中文名 F显示 G可调 H来源 I参数ID J默认 K单位 L类型
    last_sheet = ""
    print("序号|类别|所属sheet|单元格|参数中文名|F显示|G可调|数据来源|参数ID|默认值|单位|类型")
    print("-" * 8)
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=12, values_only=True):
        seq, cat, sht, cell, name, f, g, src, pid, dft, unit, dtype = row
        if seq is None and name is None:
            continue
        if sht:
            last_sheet = sht
        sht_use = sht or last_sheet
        def s(v):
            return "" if v is None else str(v).replace("\n", " ").strip()
        print(f"{s(seq)}|{s(cat)}|{s(sht_use)}|{s(cell)}|{s(name)}|{s(f)}|{s(g)}|{s(src)}|{s(pid)}|{s(dft)}|{s(unit)}|{s(dtype)}")
    wb.close()


def dump_table2_header():
    wb = openpyxl.load_workbook(TBL2, read_only=True, data_only=True)
    ws = next(w for w in wb.worksheets if w.title == "各省收益结构")
    print("\n===== 表2「各省收益结构」前 6 行 × 全列 =====")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6, max_col=19, values_only=True), start=1):
        cells = [("" if v is None else str(v)).replace("\n", " ").strip()[:14] for v in row]
        print(f"r{i:>2}: " + " | ".join(cells))
    wb.close()


if __name__ == "__main__":
    dump_mapping()
    dump_table2_header()
