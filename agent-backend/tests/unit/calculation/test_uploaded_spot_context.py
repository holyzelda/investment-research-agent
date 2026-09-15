from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.core.exceptions import ValidationException
from app.modules.calculation.infrastructure.uploaded_spot_context import (
    parse_uploaded_spot_file,
)


def _workbook_bytes(
    *, complete_days=1, bad_header=False, incomplete_day=False, gap_day=False,
    instruction_sheet=False, unexpected_sheet=False,
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "现货电价"
    sheet.append(["错误"] if bad_header else ["日期", "时点", "电价（元/MWh）"])
    for day_offset in range(complete_days):
        business_date = date(2026, 1, day_offset + 1)
        if gap_day and day_offset == 1:
            continue
        slots = 95 if incomplete_day and day_offset == complete_days - 1 else 96
        for index in range(1, slots + 1):
            minutes = index * 15
            label = f"{minutes // 60:02d}:{minutes % 60:02d}"
            price = 100 if index <= 48 else 500
            sheet.append([business_date, label, price])
    if instruction_sheet:
        workbook.create_sheet("填写说明").append(["填写说明"])
    if unexpected_sheet:
        workbook.create_sheet("其他数据").append(["其他数据"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_parses_complete_days_and_preserves_raw_curve():
    parsed = parse_uploaded_spot_file(
        _workbook_bytes(complete_days=1), "spot.xlsx", workers=1
    )

    assert parsed.valid_day_count == 1
    assert parsed.skipped_day_count == 0
    assert parsed.start_date == parsed.end_date == date(2026, 1, 1)
    assert parsed.k_results == {}
    assert parsed.price_mode == "both"
    assert len(parsed.raw_prices["2026-01-01"]) == 96
    assert parsed.raw_prices["2026-01-01"][:48] == ["100"] * 48
    assert parsed.raw_prices["2026-01-01"][48:] == ["500"] * 48


def test_skips_incomplete_day_but_reports_warning():
    parsed = parse_uploaded_spot_file(
        _workbook_bytes(complete_days=2, incomplete_day=True),
        "spot.xlsx",
        workers=1,
    )

    assert parsed.valid_day_count == 1
    assert parsed.skipped_day_count == 1
    assert "1个自然日" in parsed.warnings[0]


def test_uses_process_pool_for_multiple_complete_days():
    parsed = parse_uploaded_spot_file(
        _workbook_bytes(complete_days=2), "spot.xlsx", workers=2
    )

    assert parsed.valid_day_count == 2
    assert len(parsed.raw_prices) == 2
    assert all(len(values) == 96 for values in parsed.raw_prices.values())


def test_accepts_the_template_instruction_sheet_but_rejects_other_sheets():
    parsed = parse_uploaded_spot_file(
        _workbook_bytes(instruction_sheet=True), "spot.xlsx", workers=1
    )

    assert parsed.valid_day_count == 1
    with pytest.raises(ValidationException, match="不能包含其他工作表"):
        parse_uploaded_spot_file(
            _workbook_bytes(instruction_sheet=True, unexpected_sheet=True),
            "spot.xlsx",
            workers=1,
        )


def test_reports_validation_without_upload_time_dp_stage():
    progress = []
    parse_uploaded_spot_file(
        _workbook_bytes(complete_days=2),
        "spot.xlsx",
        workers=1,
        progress_callback=progress.append,
    )

    assert progress[0] == {"phase": "validating", "message": "正在读取并校验 Excel 数据"}
    assert progress == [
        {"phase": "validating", "message": "正在读取并校验 Excel 数据"},
    ]


def test_rejects_unknown_price_mode():
    with pytest.raises(ValidationException, match="电价用途无效"):
        parse_uploaded_spot_file(
            _workbook_bytes(complete_days=1), "spot.xlsx", price_mode="invalid"
        )


def test_counts_missing_calendar_days_in_skipped_days():
    parsed = parse_uploaded_spot_file(
        _workbook_bytes(complete_days=3, gap_day=True), "spot.xlsx", workers=1
    )

    assert parsed.valid_day_count == 2
    assert parsed.skipped_day_count == 1


def test_rejects_bad_header_and_file_type():
    with pytest.raises(ValidationException, match="表头不匹配"):
        parse_uploaded_spot_file(
            _workbook_bytes(bad_header=True), "spot.xlsx", workers=1
        )
    with pytest.raises(ValidationException, match="仅支持 .xlsx"):
        parse_uploaded_spot_file(b"x", "spot.csv", workers=1)
